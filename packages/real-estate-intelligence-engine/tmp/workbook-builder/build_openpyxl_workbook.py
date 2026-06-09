from __future__ import annotations

import csv
import json
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path("/Users/saurabhkumarbajpaiai/Documents/Codex/2026-05-06/jameenwallah-pr-real/packages/real-estate-intelligence-engine")
BASE = ROOT / "data/professional-openable-builder-db"
OUT = BASE / "JameenWallah_Builder_DB.xlsx"


def clean_cell(value: object, max_len: int = 32000) -> object:
    if value is None:
        return ""
    if not isinstance(value, str):
        return value
    cleaned = "".join(ch for ch in value.replace("\x00", "").strip() if ch == "\n" or ch == "\t" or ord(ch) >= 32)
    if len(cleaned) > max_len:
        return cleaned[: max_len - 24] + " [truncated for Excel]"
    return cleaned


def read_csv(name: str) -> list[list[object]]:
    with (BASE / name).open("r", encoding="utf-8", newline="") as handle:
        return [[clean_cell(cell) for cell in row] for row in csv.reader(handle)]


def set_tab(ws, rows: list[list[object]]) -> None:
    for row in rows:
        ws.append(row)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    header_fill = PatternFill("solid", fgColor="111827")
    header_font = Font(name="Aptos", size=10, color="FFFFFF", bold=True)
    body_font = Font(name="Aptos", size=10, color="111827")
    thin = Side(style="thin", color="E5E7EB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
            if isinstance(cell.value, str) and cell.value.startswith(("http://", "https://")):
                cell.hyperlink = cell.value
                cell.style = "Hyperlink"

    headers = [str(cell.value or "") for cell in ws[1]]
    for idx, header in enumerate(headers, 1):
        width = 18
        if any(key in header.lower() for key in ["description", "amenities", "image", "url", "address", "note", "reason", "schema", "metadata"]):
            width = 42
        elif any(key in header.lower() for key in ["project_name", "builder_name", "locality", "city", "sector"]):
            width = 26
        ws.column_dimensions[get_column_letter(idx)].width = width


def add_readme(wb: Workbook, summary: dict) -> None:
    ws = wb.active
    ws.title = "README"
    rows = [
        ["JameenWallah Builder Database", ""],
        ["Generated", datetime.now(timezone.utc).isoformat()],
        ["Best import sheet", "Production Ready"],
        ["Full curated sheet", "Properties Curated"],
        ["Image relationship sheet", "Image Slots"],
        ["Production caution", "Rows in Review Needed should be checked before direct production import."],
        ["", ""],
        ["Metric", "Value"],
        ["Valid property rows", summary["rows_valid_properties"]],
        ["Curated rows", summary["rows_curated"]],
        ["Image slot rows", summary["rows_image_slots"]],
        ["Excluded non-property rows", summary["rows_excluded"]],
        ["Rows needing review", summary["rows_review_needed"]],
        ["Rows production-ready with 5 clean images", summary.get("rows_production_ready_with_5_clean_images", "")],
        ["Rows needing manual image sourcing", summary.get("rows_manual_image_sourcing_needed", "")],
        ["Sanitized corrupted row IDs", ", ".join(summary.get("sanitized_corrupted_row_ids", []))],
        ["", ""],
        ["Local files", ""],
        ["CSV import file", "properties_curated.csv"],
        ["All image slots", "image_slots_openable.csv"],
        ["SQLite database", "professional_builder_db.sqlite"],
        ["Schema", "schema.json / SCHEMA.md"],
    ]
    set_tab(ws, rows)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="0F766E")
        cell.font = Font(name="Aptos", size=14, color="FFFFFF", bold=True)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 84


def add_schema(wb: Workbook, schema: dict) -> None:
    rows = [["table", "field", "type", "description"]]
    for table_name, table in schema.get("tables", {}).items():
        fields = table.get("fields")
        if fields:
            iterable = [(field.get("name", ""), field.get("type", ""), field.get("description", "")) for field in fields]
        else:
            iterable = [
                (field_name, meta.get("type", ""), meta.get("description", ""))
                for field_name, meta in table.get("columns", {}).items()
            ]
        for field_name, field_type, description in iterable:
            rows.append([
                table_name,
                field_name,
                field_type,
                description,
            ])
    ws = wb.create_sheet("Schema")
    set_tab(ws, rows)


def main() -> None:
    summary = json.loads((BASE / "summary.json").read_text(encoding="utf-8"))
    schema = json.loads((BASE / "schema.json").read_text(encoding="utf-8"))

    wb = Workbook()
    add_readme(wb, summary)

    for title, filename in [
        ("Production Ready", "production_ready_properties.csv"),
        ("Properties Curated", "properties_curated.csv"),
        ("Image Slots", "image_slots_openable.csv"),
        ("Manual Image Sourcing", "manual_image_sourcing_needed.csv"),
        ("Review Needed", "review_needed.csv"),
        ("Excluded Rows", "excluded_non_property_rows.csv"),
    ]:
        ws = wb.create_sheet(title)
        set_tab(ws, read_csv(filename))

    add_schema(wb, schema)
    wb.save(OUT)

    check = load_workbook(OUT, read_only=True, data_only=True)
    print(f"saved={OUT}")
    for ws in check.worksheets:
        print(f"{ws.title}: rows={ws.max_row} cols={ws.max_column}")


if __name__ == "__main__":
    main()
